"""
Scraper CA · Alberta Drug Benefit List (ADBL / Blue Cross).

Excel público mensual con precio de beneficio (CAD) por DIN.
No es retail de farmacia online; es listado oficial provincial.
"""

from __future__ import annotations

import io
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

import requests
import urllib3
from dotenv import load_dotenv
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
load_dotenv(ROOT / ".env")

from lista import MEDICAMENTOS  # noqa: E402
from scrapper.ficha import parse_ficha  # noqa: E402
from scrapper.matching import clasificar, coincide  # noqa: E402
from scrapper.report_seeds import seed_match_for_med  # noqa: E402
from scrapper.repositorio import asegurar_tabla, contar, guardar_filas, purgar_obsoletos  # noqa: E402

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

CACHE = ROOT / "cache" / "farmacias" / "adbl_ca"
PUB_URL = "https://www.ab.bluecross.ca/dbl/publications.php"
BASE = "https://www.ab.bluecross.ca"
PAIS = "Canadá"
FARMACIA = "ADBL Alberta"
MONEDA = "CAD"
TIMEOUT = 120
UA = (
    "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
)

MESES = {
    "jan": 1,
    "feb": 2,
    "mar": 3,
    "apr": 4,
    "may": 5,
    "jun": 6,
    "jul": 7,
    "aug": 8,
    "sep": 9,
    "oct": 10,
    "nov": 11,
    "dec": 12,
}


def session() -> requests.Session:
    s = requests.Session()
    s.verify = False
    s.headers.update({"User-Agent": UA, "Accept": "*/*"})
    return s


def parse_money(raw: Any) -> float | None:
    if raw is None:
        return None
    s = str(raw).strip().replace(",", "")
    s = re.sub(r"[^0-9.]", "", s)
    if not s:
        return None
    try:
        n = float(s)
    except ValueError:
        return None
    return n if n > 0 else None


def elegir_xlsx(html: str) -> str:
    """Prefiere dbl_<mes>.xlsx (listado mensual con PRICE)."""
    links = re.findall(r'href=["\']([^"\']*dbl_[a-z]{3}\.xlsx)["\']', html, re.I)
    if not links:
        links = re.findall(r'href=["\']([^"\']*ADBL_HSDBS[^"\']*\.xlsx)["\']', html, re.I)
    scored: list[tuple[int, str]] = []
    for href in links:
        m = re.search(r"dbl_([a-z]{3})\.xlsx", href, re.I)
        mes = MESES.get((m.group(1).lower() if m else ""), 0)
        scored.append((mes, href))
    if not scored:
        raise RuntimeError("No se encontró Excel ADBL en publications.php")
    scored.sort(key=lambda x: x[0], reverse=True)
    href = scored[0][1]
    if href.startswith("http"):
        return href
    return BASE.rstrip("/") + "/" + href.lstrip("/")


def descargar_catalogo(s: requests.Session, *, forzar: bool) -> tuple[list[dict[str, Any]], str]:
    CACHE.mkdir(parents=True, exist_ok=True)
    meta_path = CACHE / "_source.txt"
    xlsx_path = CACHE / "adbl.xlsx"
    if not forzar and xlsx_path.exists() and xlsx_path.stat().st_size > 50_000:
        url = meta_path.read_text(encoding="utf-8").strip() if meta_path.exists() else ""
        return leer_xlsx(xlsx_path.read_bytes()), url or str(xlsx_path)

    r = s.get(PUB_URL, timeout=TIMEOUT)
    r.raise_for_status()
    url = elegir_xlsx(r.text)
    print(f"  bajando {url} …")
    rx = s.get(url, timeout=TIMEOUT)
    rx.raise_for_status()
    xlsx_path.write_bytes(rx.content)
    meta_path.write_text(url, encoding="utf-8")
    return leer_xlsx(rx.content), url


def _norm_header(h: Any) -> str:
    s = str(h or "").strip().upper()
    s = re.sub(r"\s+", " ", s)
    return s


def leer_xlsx(raw: bytes) -> list[dict[str, Any]]:
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    out: list[dict[str, Any]] = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        hdr: list[str] | None = None
        for row in ws.iter_rows(values_only=True):
            if hdr is None:
                hdr = [_norm_header(c) for c in row]
                continue
            d = {hdr[i]: row[i] if i < len(row) else None for i in range(len(hdr))}
            # Filas espaciadoras del Excel (DIN numérico suelto)
            din = d.get("DIN / PIN") or d.get("DIN/ NPN /PIN") or d.get("DIN/NPN/PIN")
            if din is None or (isinstance(din, (int, float)) and float(din) < 1000):
                if not d.get("PRODUCT NAME") and not d.get("GENERIC NAME"):
                    continue
            precio = parse_money(d.get("PRICE")) or parse_money(d.get("LCA / MAC PRICE"))
            if not precio:
                continue
            nombre = str(d.get("PRODUCT NAME") or "").strip()
            generico = str(d.get("GENERIC NAME") or "").strip()
            if not nombre and not generico:
                continue
            d["_sheet"] = sheet
            d["_price"] = precio
            d["_din"] = str(din or "").strip()
            out.append(d)
    return out


def pack_units(row: dict[str, Any]) -> int:
    """ADBL publica precio por unidad (tab/cáps/mL/dispositivo)."""
    name = str(row.get("PRODUCT NAME") or "")
    form = str(row.get("FORM") or "").upper()
    route = str(row.get("ROUTE") or "").upper()
    blob = f"{name} {form} {route}".upper()

    m = re.search(r"\((\d+(?:\.\d+)?)\s*ML\)", name, re.I)
    if m and ("INJ" in blob or "ML" in blob):
        vol = float(m.group(1))
        if vol >= 1:
            return int(round(vol))

    if re.search(r"(SYRINGE|PEN|AUTO\s*INJECTOR|AUTOINJECTOR|/VIAL| KIT\b)", name, re.I):
        return 1

    if any(x in blob for x in ("TAB", "CAP", "ORAL")) and "INJ" not in blob:
        return 30
    return 1


def filas_de_catalogo(
    med: dict[str, Any],
    productos: list[dict[str, Any]],
    fecha_dato: str,
    fuente: str,
    vistos: set[tuple[str, int]],
) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for row in productos:
        nombre = str(row.get("PRODUCT NAME") or "").strip()
        generico = str(row.get("GENERIC NAME") or "").strip()
        fuerza = str(row.get("STRENGTH") or "").strip()
        forma = str(row.get("FORM") or "").strip()
        blob = " ".join(x for x in (nombre, generico, fuerza, forma) if x)
        seed_match = seed_match_for_med(blob, PAIS, med)
        if not coincide(blob, med) and not seed_match:
            continue
        din = str(row.get("_din") or "").strip() or nombre[:40]
        n_lista = int(med["n"])
        clave = (din, n_lista)
        if clave in vistos:
            continue
        unit = float(row["_price"])
        units = pack_units(row)
        precio = round(unit * units, 2)
        if precio <= 0:
            continue
        vistos.add(clave)
        label = " ".join(x for x in (nombre, fuerza) if x) or generico
        p_fake = {
            "nombre": label,
            "principios_activos": [generico, nombre],
            "tipo_presentacion": f"{forma or 'unit'} x {units}",
            "laboratorio": str(row.get("MFR") or "") or None,
        }
        med2, calidad, obs = clasificar(med, p_fake)
        if seed_match and calidad == "ok":
            calidad = "revisar"
        note = f"ADBL unitario CAD {unit} × {units} ({row.get('_sheet')})"
        obs = f"{obs}; {note}" if obs else note
        ficha = parse_ficha(label, med2)
        out.append(
            {
                "pais": PAIS,
                "farmacia": FARMACIA,
                "fuente_url": (fuente or PUB_URL)[:500],
                "id_producto_farmacia": din[:80],
                "sku": din[:80],
                "n_lista": int(med2["n"]),
                "medicamento_lista": str(med2["nombre"]),
                "programa": str(med2["programa"]),
                "nombre_comercial": label[:500],
                "principio_activo": ficha.get("principio_activo") or generico[:200] or None,
                "concentracion": ficha.get("concentracion") or (fuerza[:80] if fuerza else None),
                "presentacion": ficha.get("presentacion")
                or f"{forma or 'unit'} × {units}".strip()[:120],
                "laboratorio": (str(row.get("MFR") or "").strip() or ficha.get("laboratorio") or None),
                "precio": precio,
                "moneda": MONEDA,
                "disponibilidad": str(row.get("COVERAGE STATUS") or "Listado")[:80],
                "calidad": calidad,
                "observacion": obs,
                "fecha_publicacion": None,
                "fecha_dato": fecha_dato,
            }
        )
    return out


def recolectar(*, forzar: bool = False, meds: list[dict[str, Any]] | None = None) -> list[dict[str, Any]]:
    s = session()
    print("  bajando Alberta Drug Benefit List…")
    productos, fuente = descargar_catalogo(s, forzar=forzar)
    print(f"  catálogo: {len(productos)} filas con precio · {fuente}")
    fecha = datetime.now().date().isoformat()
    filas: list[dict[str, Any]] = []
    vistos: set[tuple[str, int]] = set()
    trabajo = meds if meds is not None else MEDICAMENTOS
    for med in trabajo:
        n_antes = len(filas)
        filas.extend(filas_de_catalogo(med, productos, fecha, fuente, vistos))
        n_nuevas = len(filas) - n_antes
        marca = "·" if n_nuevas else " "
        print(f"  {marca} {n_nuevas:3d}  {med['nombre']}")
    return filas


def main() -> int:
    forzar = "--fresh" in sys.argv
    solo_fomac = "--fomac" in sys.argv
    meds = MEDICAMENTOS
    if solo_fomac:
        meds = [m for m in MEDICAMENTOS if "FOMAC" in str(m["programa"])]
    print("=" * 64)
    print(" CA ADBL Alberta → medicamentos_altos_costos_america")
    if solo_fomac:
        print(f" Solo FOMAC ({len(meds)} moléculas)")
    print("=" * 64)
    filas = recolectar(forzar=forzar, meds=meds)
    ok = sum(1 for f in filas if f["calidad"] == "ok")
    print(f"Lista: {len(meds)} medicamentos")
    print(f"Coincidencias a guardar: {len(filas)}  (ok={ok}, revisar={len(filas) - ok})")
    asegurar_tabla()
    res = guardar_filas(filas)
    borrados = 0
    if filas:
        borrados = purgar_obsoletos(
            PAIS,
            FARMACIA,
            [(f["id_producto_farmacia"], f["n_lista"]) for f in filas],
        )
    else:
        print("  aviso: 0 coincidencias — no se purgan filas existentes")
    print(
        f"MERGE {res['upserts']} filas · obsoletos: {borrados} · "
        f"{PAIS} {FARMACIA}: {contar(PAIS, FARMACIA)} · total tabla: {contar()}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
